"""Export the Tier 3 comparison chart as editable SVG and PowerPoint shapes."""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path("/workspace")
SUMMARY = ROOT / "pems_tier3_co2_calibration_summary.csv"
SVG_OUT = ROOT / "pems_tier3_co2_calibration_editable.svg"
PPTX_OUT = ROOT / "pems_tier3_co2_calibration_editable.pptx"
ROADS = ["快速路", "主干路", "次干路", "支路"]


def load_summary():
    table = pd.read_csv(SUMMARY)
    return table[table["道路类型"].isin(ROADS)].reset_index(drop=True)


def export_svg(table):
    font_path = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    font_manager.fontManager.addfont(font_path)
    plt.rcParams.update({"font.family": font_manager.FontProperties(fname=font_path).get_name(),
                         "axes.unicode_minus": False, "font.size": 10})
    fig, axes = plt.subplots(1, 2, figsize=(14, 6.2), constrained_layout=True)
    x = range(len(table))
    width = 0.34
    axes[0].bar([v - width / 2 for v in x], table["实测值（g/km）"], width,
                label="PEMS实测值", color="#4E79A7")
    axes[0].bar([v + width / 2 for v in x], table["COPERT模拟值（g/km）"], width,
                label="COPERT模拟值", color="#CC5A6A")
    axes[0].set_xticks(list(x), ROADS)
    axes[0].set_ylabel("CO2排放因子（g/km）")
    axes[0].set_title("四种道路类型的实测值与模拟值")
    axes[0].grid(axis="y", color="#DCEAF6", linewidth=0.8)
    axes[0].legend()
    axes[1].bar(list(x), table["COPERT误差率（%）"], width * 1.25,
                label="COPERT模拟误差率", color="#CC5A6A")
    axes[1].set_xticks(list(x), ROADS)
    axes[1].set_ylabel("误差率（%）")
    axes[1].set_title("COPERT模拟值与实测值的误差率")
    axes[1].grid(axis="y", color="#DCEAF6", linewidth=0.8)
    axes[1].legend()
    fig.suptitle("Tier 3 CO2排放因子与PEMS实测值对比", fontsize=17, weight="bold")
    fig.savefig(SVG_OUT, format="svg", bbox_inches="tight")
    plt.close(fig)


def add_text(slide, text, left, top, width, height, size=12, bold=False,
             color=(30, 30, 30), align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_line(slide, x1, y1, x2, y2, color=(220, 234, 246), width=1):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor(*color)
    line.line.width = Pt(width)
    return line


def add_panel(slide, table, left, top, width, height, kind):
    plot_left, plot_top = left + 0.58, top + 0.58
    plot_width, plot_height = width - 0.82, height - 1.25
    if kind == "comparison":
        title = "四种道路类型的实测值与模拟值"
        ylabel = "CO2排放因子（g/km）"
        max_value = max(table["实测值（g/km）"].max(), table["COPERT模拟值（g/km）"].max()) * 1.18
        series = [("实测值", table["实测值（g/km）"], (78, 121, 167)),
                  ("COPERT模拟值", table["COPERT模拟值（g/km）"], (204, 90, 106))]
    else:
        title = "COPERT模拟值与实测值的误差率"
        ylabel = "误差率（%）"
        max_value = max(table["COPERT误差率（%）"].max() * 1.25, 5)
        series = [("COPERT模拟误差率", table["COPERT误差率（%）"], (204, 90, 106))]
    add_text(slide, title, left, top, width, 0.42, size=14, bold=False)
    add_text(slide, ylabel, left, top + 0.42, 0.48, 0.38, size=9)
    baseline = plot_top + plot_height
    for tick in range(0, int(max_value) + 1, max(1, int(max_value / 5))):
        y = baseline - plot_height * tick / max_value
        add_line(slide, plot_left, y, plot_left + plot_width, y)
        add_text(slide, tick, left, y - 0.12, 0.48, 0.22, size=8, align=PP_ALIGN.RIGHT)
    add_line(slide, plot_left, plot_top, plot_left, baseline, color=(80, 80, 80), width=1.2)
    add_line(slide, plot_left, baseline, plot_left + plot_width, baseline, color=(80, 80, 80), width=1.2)
    group_width = plot_width / len(ROADS)
    bar_width = min(0.24, group_width / 4)
    for i, road in enumerate(ROADS):
        center = plot_left + group_width * (i + 0.5)
        add_text(slide, road, center - 0.45, baseline + 0.06, 0.9, 0.25, size=9)
        for j, (_, values, color) in enumerate(series):
            if kind == "comparison":
                x = center + (j - 0.5) * bar_width
            else:
                x = center
            value = float(values.iloc[i])
            bar_height = plot_height * value / max_value
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - bar_width / 2),
                                           Inches(baseline - bar_height), Inches(bar_width), Inches(bar_height))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*color)
            shape.line.color.rgb = RGBColor(*color)
            add_text(slide, f"{value:.1f}", x - 0.35, baseline - bar_height - 0.24, 0.7, 0.22, size=8)
    legend_top = top + height - 0.36
    legend_left = left + 0.7
    for j, (label, _, color) in enumerate(series):
        x = legend_left + j * 1.35
        swatch = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(legend_top + 0.04), Inches(0.16), Inches(0.16))
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = RGBColor(*color)
        swatch.line.color.rgb = RGBColor(*color)
        add_text(slide, label, x + 0.2, legend_top, 1.1, 0.25, size=8, align=PP_ALIGN.LEFT)


def export_pptx(table):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "Tier 3 CO2排放因子与PEMS实测值对比", 0.3, 0.12, 12.7, 0.42, size=20, bold=True)
    add_panel(slide, table, 0.35, 0.75, 6.2, 6.25, "comparison")
    add_panel(slide, table, 6.8, 0.75, 6.2, 6.25, "error")
    prs.save(PPTX_OUT)


def main():
    table = load_summary()
    export_svg(table)
    export_pptx(table)
    print(SVG_OUT)
    print(PPTX_OUT)


if __name__ == "__main__":
    main()
